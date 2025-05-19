#adminchat/services/embedding_service.py
import re
import requests
from django.conf import settings
from urllib.parse import urljoin
import boto3
from botocore.exceptions import ClientError
import logging
from io import BytesIO
from PyPDF2 import PdfReader
from docx import Document as DocxDocument
import openpyxl
import csv

logger = logging.getLogger(__name__)

class TextExtractor:
    """Clase para extraer texto de diferentes tipos de archivos"""
    
    @staticmethod
    def extract_from_pdf(file_content):
        """Extrae texto de un PDF"""
        try:
            pdf_reader = PdfReader(BytesIO(file_content))
            text = ""
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
            return text
        except Exception as e:
            logger.error(f"Error extracting PDF text: {str(e)}")
            raise ValueError("Could not extract text from PDF file")

    @staticmethod
    def extract_from_docx(file_content):
        """Extrae texto de un DOCX"""
        try:
            doc = DocxDocument(BytesIO(file_content))
            return "\n".join([para.text for para in doc.paragraphs])
        except Exception as e:
            logger.error(f"Error extracting DOCX text: {str(e)}")
            raise ValueError("Could not extract text from Word file")

    @staticmethod
    def extract_from_txt(file_content):
        """Extrae texto de un TXT"""
        try:
            return file_content.decode('utf-8')
        except Exception as e:
            logger.error(f"Error extracting TXT text: {str(e)}")
            raise ValueError("Could not extract text from text file")

    @staticmethod
    def extract_from_xlsx(file_content):
        """Extrae texto de un XLSX"""
        try:
            wb = openpyxl.load_workbook(BytesIO(file_content))
            text = ""
            for sheet in wb.worksheets:
                for row in sheet.iter_rows():
                    text += " ".join([str(cell.value) for cell in row if cell.value]) + "\n"
            return text
        except Exception as e:
            logger.error(f"Error extracting XLSX text: {str(e)}")
            raise ValueError("Could not extract text from Excel file")

    @staticmethod
    def extract_from_csv(file_content):
        """Extrae texto de un CSV"""
        try:
            text = ""
            reader = csv.reader(BytesIO(file_content).read().decode('utf-8').splitlines())
            for row in reader:
                text += " ".join(row) + "\n"
            return text
        except Exception as e:
            logger.error(f"Error extracting CSV text: {str(e)}")
            raise ValueError("Could not extract text from CSV file")

class TextCleaner:
    """Clase para limpieza de texto"""
    
    @staticmethod
    def clean_text(text: str) -> str:
        """Limpia el texto según las reglas especificadas"""
        text = text.replace('\r\n', '\n').replace('\r', '\n')  # normaliza saltos de línea
        text = re.sub(r'[ \t]+', ' ', text)                   # reemplaza múltiples espacios por uno
        text = re.sub(r'\n{2,}', '\n\n', text)                # reduce saltos de línea excesivos
        return text.strip()

class ChunkGenerator:
    """Clase para generar chunks de texto"""
    
    @staticmethod
    def generate_chunks(text: str, chunk_size: int, chunk_overlap: int) -> list:
        """Divide el texto en chunks según tamaño y overlap especificado"""
        words = text.split()
        chunks = []
        
        if chunk_size <= 0:
            raise ValueError("chunk_size must be greater than 0")
        
        if chunk_overlap >= chunk_size:
            raise ValueError("chunk_overlap must be smaller than chunk_size")
        
        for i in range(0, len(words), chunk_size - chunk_overlap):
            chunk = words[i:i + chunk_size]
            chunks.append(" ".join(chunk))
        
        return chunks

class S3FileService:
    """Clase para interactuar con S3"""
    
    def __init__(self):
        self.s3_client = boto3.client(
            's3',
            aws_access_key_id=settings.AWS_ACCESS_KEY_ID,
            aws_secret_access_key=settings.AWS_SECRET_ACCESS_KEY,
            region_name=settings.AWS_REGION
        )
        self.bucket_name = settings.S3_BUCKET
    
    def get_file_content(self, s3_path: str) -> bytes:
        """Obtiene el contenido de un archivo desde S3"""
        if not s3_path.startswith('s3://'):
            raise ValueError("Invalid S3 path format")
        
        # Extraer bucket y key del path
        path_parts = s3_path[5:].split('/', 1)  # Remueve 's3://'
        bucket_name = path_parts[0]
        object_key = path_parts[1] if len(path_parts) > 1 else ''
        
        try:
            response = self.s3_client.get_object(Bucket=bucket_name, Key=object_key)
            return response['Body'].read()
        except ClientError as e:
            logger.error(f"Error getting file from S3: {str(e)}")
            raise ValueError(f"Could not retrieve file from S3: {str(e)}")

class EmbeddingGenerator:
    """Clase para generar embeddings llamando al servicio externo"""
    
    @staticmethod
    def generate_embeddings(texts: list, embedding_model: str, embedding_dim: int) -> list:
        """Llama al servicio de embeddings para vectorizar los textos"""
        print(f'embedding_model: {embedding_model}')
        print(f'embedding_dim: {embedding_dim}')

        try:
            # URL del servicio de embeddings (ajustar según configuración)
            url=settings.URL_EMBEDDING
            #embedding_service_url = "http://embedder.local:8000/api/v1/embeddings/generate"
            embedding_service_url = F"{url}/api/v1/embeddings/generate"
            payload = {
                "texts": texts, 
                "embedding_model": embedding_model,
                "embedding_dim": embedding_dim
            }
            
            response = requests.post(embedding_service_url, json=payload)
            response.raise_for_status()
            
            return response.json().get('embeddings', [])
        except requests.exceptions.RequestException as e:
            logger.error(f"Error calling embedding service: {str(e)}")
            raise ValueError(f"Could not generate embeddings: {str(e)}")